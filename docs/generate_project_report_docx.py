"""Generate a structured Word version of the SPARC-PM project report."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "project_report.md"
OUTPUT = ROOT / "docs" / "SPARC_PM_Project_Report.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_toc(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    paragraph._p.append(field)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.75)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 25, "17365D"),
        ("Heading 1", 17, "17365D"),
        ("Heading 2", 13, "1F4E79"),
        ("Heading 3", 11, "2F5597"),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code Listing" not in [style.name for style in styles]:
        code_style = styles.add_style("Code Listing", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Listing"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.2)
    code_style.font.color.rgb = RGBColor(40, 40, 40)
    code_style.paragraph_format.left_indent = Inches(0.2)
    code_style.paragraph_format.space_after = Pt(0)

    header = section.header.paragraphs[0]
    header.text = "SPARC-PM | Project Report"
    header.style = styles["Header"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SPARC-PM")
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = RGBColor(23, 54, 93)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Space Priority Alert & Response Command Engine\n"
        "with Predictive Maintenance"
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(47, 85, 151)

    document.add_paragraph()
    meta = document.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    fields = [
        ("Project domain", "Autonomous mission operations and space-weather survivability"),
        ("Target operations centre", "ISRO Mission Control Centre simulation"),
        ("Repository", "https://github.com/RomilaMukul/SPARC"),
        ("Submission version", "Version 2.4.0 (production release)"),
        ("Student / institution", "[Complete before submission]"),
    ]
    for row, (label, value) in zip(meta.rows, fields):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "D9EAF7")
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        row.cells[0].paragraphs[0].runs[0].bold = True
    document.add_page_break()

    toc_heading = document.add_paragraph("Table of contents", style="Heading 1")
    toc_heading.paragraph_format.page_break_before = False
    toc = document.add_paragraph()
    add_toc(toc)
    document.add_paragraph(
        "Note: In Microsoft Word, right-click the table of contents and choose "
        "'Update Field' to refresh page numbers after editing."
    ).italic = True
    document.add_page_break()


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) < 2:
        return
    header = rows[0]
    body = [row for row in rows[2:] if not all(set(part) <= {"-", ":", " "} for part in row)]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for index, value in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, "17365D")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
    for row_values in body:
        row = table.add_row()
        for index, value in enumerate(row_values[: len(header)]):
            cell = row.cells[index]
            cell.text = value
            set_cell_margins(cell)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cell, "F2F6FA")
    document.add_paragraph()


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Listing")
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.1)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        if index < len(lines) - 1:
            run.add_break()
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F3F3")
    properties.append(shading)


def add_text(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    clean = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    clean = clean.replace("**", "").replace("*", "")
    paragraph.add_run(clean)


def convert_markdown(document: Document, source: str) -> None:
    lines = source.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).replace(r"\*", "*")
            document.add_heading(title, level=min(level, 3))
            index += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", bullet.group(1)))
            index += 1
            continue
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(numbered.group(1))
            index += 1
            continue
        add_text(document, line)
        index += 1


def main() -> None:
    document = Document()
    configure_document(document)
    add_cover(document)
    source = SOURCE.read_text(encoding="utf-8")
    convert_markdown(document, source)
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
